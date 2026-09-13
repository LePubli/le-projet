"""F: Notifications email — envoi d'emails sur pannes critiques.
I: Alertes stock — seuil d'alerte sur les pièces.
D: Export Excel natif (.xlsx).
E: Chat/messagerie interne sur pannes.
J: Signature électronique sur BT.

Ce module regroupe les nouveaux endpoints pour les fonctionnalités C à J.
"""

import os
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy.orm import Session
from pydantic import BaseModel
from app.db.database import get_db
from app.models.panne import Panne
from app.models.bon_travail import BonTravail
from app.core.security import (
    get_current_user,
    require_manager_or_admin,
    require_edit_access,
    require_signatory_role,
)
from app.core.logging import get_logger
from app.services.extra import (
    append_chat_message,
    build_signature_marker,
    get_stock_alertes,
    get_xlsx_data,
    parse_chat_messages,
    parse_signatures,
)

log = get_logger(__name__)

router = APIRouter(prefix="/api/extra", tags=["extra"])


# ═══════════════════════════════════════════════════════════════
#  F — NOTIFICATIONS EMAIL
# ═══════════════════════════════════════════════════════════════


class EmailConfig(BaseModel):
    smtp_host: str = ""
    smtp_port: int = 587
    smtp_user: str = ""
    smtp_password: str = ""
    from_email: str = ""
    enabled: bool = False


@router.get("/email/status")
def email_status(_=Depends(get_current_user)):
    """Vérifie si l'email est configuré."""
    return {"enabled": bool(os.getenv("SMTP_HOST", ""))}


# ═══════════════════════════════════════════════════════════════
#  I — ALERTES STOCK
# ═══════════════════════════════════════════════════════════════


@router.get("/stock/alertes")
def stock_alertes(
    db: Session = Depends(get_db),
    _=Depends(get_current_user),
):
    """Liste les pièces dont le stock est en dessous du seuil d'alerte."""
    return get_stock_alertes(db)


# ═══════════════════════════════════════════════════════════════
#  D — EXPORT EXCEL NATIF (.xlsx)
# ═══════════════════════════════════════════════════════════════


@router.get("/export/xlsx/{entity}")
def export_xlsx(
    entity: str,
    db: Session = Depends(get_db),
    _=Depends(require_manager_or_admin),
):
    """Exporte une entité au format Excel .xlsx avec mise en forme."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from fastapi.responses import StreamingResponse

    try:
        headers, rows = get_xlsx_data(entity, db)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    wb = Workbook()
    ws = wb.active
    ws.title = entity

    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="F97316", end_color="F97316", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    ws.append(headers)
    for row in rows:
        ws.append(row)

    # Mise en forme en-têtes
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Auto-largeur colonnes
    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 40)

    # Filtres auto
    ws.auto_filter.ref = ws.dimensions

    # Freeze première ligne
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"export_{entity}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ═══════════════════════════════════════════════════════════════
#  E — CHAT/MESSAGERIE INTERNE SUR PANNES
# ═══════════════════════════════════════════════════════════════


class ChatMessage(BaseModel):
    message: str


@router.get("/chat/{panne_id}")
def get_chat(
    panne_id: int,
    db: Session = Depends(get_db),
    _=Depends(get_current_user),
):
    """Récupère les messages de chat d'une panne (stockés dans Panne.description)."""
    panne = db.query(Panne).filter(Panne.id == panne_id).first()
    if not panne:
        raise HTTPException(status_code=404, detail="Panne introuvable")
    return {"messages": parse_chat_messages(panne)}


@router.post("/chat/{panne_id}")
def post_chat(
    panne_id: int,
    body: ChatMessage,
    db: Session = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Ajoute un message de chat sur une panne."""
    panne = db.query(Panne).filter(Panne.id == panne_id).first()
    if not panne:
        raise HTTPException(status_code=404, detail="Panne introuvable")

    new_msg = append_chat_message(db, panne, current_user.username, current_user.role, body.message)
    return {"ok": True, "message": new_msg}


# ═══════════════════════════════════════════════════════════════
#  J — SIGNATURE ÉLECTRONIQUE SUR BT
# ═══════════════════════════════════════════════════════════════


class SignatureBody(BaseModel):
    signature_data: str  # base64 PNG data URL
    role: str = "intervenant"  # intervenant | responsable | controle


@router.post("/bt/{bt_id}/signature")
def add_signature(
    bt_id: int,
    body: SignatureBody,
    db: Session = Depends(get_db),
    current_user=Depends(require_edit_access),
):
    """Ajoute une signature électronique sur un Bon de Travail.
    La signature est stockée comme data URL dans le champ compte_rendu
    (préfixé [SIG:role])."""
    bt = db.query(BonTravail).filter(BonTravail.id == bt_id).first()
    if not bt:
        raise HTTPException(status_code=404, detail="BT introuvable")

    require_signatory_role(current_user.role, body.role)

    # Stocker la signature
    sig_marker = build_signature_marker(body.role, current_user.username, body.signature_data)
    if bt.compte_rendu:
        bt.compte_rendu = bt.compte_rendu + "\n" + sig_marker
    else:
        bt.compte_rendu = sig_marker
    db.commit()

    return {"ok": True, "role": body.role, "signed_by": current_user.username}


@router.get("/bt/{bt_id}/signatures")
def get_signatures(
    bt_id: int,
    db: Session = Depends(get_db),
    _=Depends(get_current_user),
):
    """Récupère les signatures d'un BT."""
    bt = db.query(BonTravail).filter(BonTravail.id == bt_id).first()
    if not bt:
        raise HTTPException(status_code=404, detail="BT introuvable")

    return {"signatures": parse_signatures(bt.compte_rendu)}
